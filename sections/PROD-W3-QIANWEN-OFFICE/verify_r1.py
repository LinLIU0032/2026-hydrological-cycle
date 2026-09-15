import sys, io, zipfile, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

path = 'E:/PG/2026 Hydrological cycle/sections/PROD-W3-QIANWEN-OFFICE/L16_W3_QianwenOffice_Slides36-37_39_v01.pptx'
prs = Presentation(path)
print('slide_count =', len(list(prs.slides)))
print('size_in = %.4f x %.4f' % (prs.slide_width/914400, prs.slide_height/914400))

def walk(sh):
    yield sh
    if sh.shape_type == 6:
        for s in sh.shapes:
            yield from walk(s)

BANNED_VISIBLE = ['7%', '75%', 'Snowball', '植物生长可促进']
for i, slide in enumerate(prs.slides, 1):
    vis = []
    pics = 0
    for sh in slide.shapes:
        for s in walk(sh):
            if s.shape_type == 13:
                pics += 1
            if s.has_text_frame:
                vis.append(s.text_frame.text)
    joined = '\n'.join(vis)
    note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    hits = [b for b in BANNED_VISIBLE if b in joined]
    nhits = [b for b in ['7%', '75%', '植物生长可促进'] if b in note]
    print(f'slide {i}: pictures={pics} notes_chars={len(note)} banned_visible={hits} banned_notes={nhits}')
    if i == 3:
        print('  slide3 has 植物生长可能增加 visible:', '植物生长可能增加' in joined, '| in notes:', '植物生长可能增加' in note)
    if i == 2:
        print('  slide2 keeps albedo definition visible:', ('反照率 Albedo = 地表反射率 surface reflectivity' in joined))
    if i == 1:
        print('  slide1 source ends with context only:', joined.strip().splitlines()[-1].endswith('supports phase-change context only.') if joined.strip().splitlines() else False)
    print('  sign:', ('Positive Feedback 正反馈' in joined), ('Negative Feedback 负反馈' in joined))

z = zipfile.ZipFile(path)
ext = sum(1 for n in z.namelist() if n.startswith('ppt/slides/_rels/') and b'TargetMode="External"' in z.read(n))
anim = audio = adv = 0
for n in z.namelist():
    if re.match(r'ppt/slides/slide\d+\.xml$', n):
        x = z.read(n).decode('utf-8')
        anim += ('<p:anim' in x or '<p:set' in x)
        audio += ('p:audio' in x or '.mp3' in x or '.wav' in x)
        adv += ('advTm' in x)
print('external =', ext, 'media =', [n for n in z.namelist() if n.startswith('ppt/media/')], 'anim =', anim, 'audio =', audio, 'timed_advance =', adv)
raw = open('E:/PG/2026 Hydrological cycle/sections/PROD-W3-QIANWEN-OFFICE/L16_W3_QianwenOffice_Slides36-37_39_v01.pdf', 'rb').read()
print('pdf_count =', re.findall(rb'/Count \d+', raw)[:2], 'page_objs =', len(re.findall(rb'/Type/Page[^s]', raw)))
print('R1_VERIFY_DONE')
