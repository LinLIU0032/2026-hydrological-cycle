import sys, io, zipfile, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

path = 'E:/PG/2026 Hydrological cycle/sections/PROD-W3-QIANWEN-OFFICE/L16_W3_QianwenOffice_Slides36-37_39_v01.pptx'
prs = Presentation(path)
print('slide_count =', len(list(prs.slides)))
print('size_in = %.4f x %.4f' % (prs.slide_width/914400, prs.slide_height/914400))

def walk(sh):
    yield sh
    if sh.shape_type == 6:
        for s in sh.shapes:
            yield from walk(s)

for i, slide in enumerate(prs.slides, 1):
    pics = 0
    total = 0
    for sh in slide.shapes:
        for s in walk(sh):
            total += 1
            if s.shape_type == 13:
                pics += 1
    note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    print(f'slide {i}: recursive_shapes={total} pictures={pics} notes_chars={len(note)}')

# terminology + sign checks on visible text
import collections
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        for s in walk(sh):
            if s.has_text_frame:
                texts.append(s.text_frame.text)
    joined = '\n'.join(texts)
    print(f'slide {i}: has_daqiquan_wrong={"大圈" in joined.replace("大气圈","")}, positive={"Positive Feedback 正反馈" in joined}, negative={"Negative Feedback 负反馈" in joined}')

z = zipfile.ZipFile(path)
ext = 0
for n in z.namelist():
    if n.startswith('ppt/slides/_rels/'):
        if b'TargetMode="External"' in z.read(n):
            ext += 1
print('external_relationships =', ext)
print('media =', [n for n in z.namelist() if n.startswith('ppt/media/')])
anim = audio = 0
for n in z.namelist():
    if re.match(r'ppt/slides/slide\d+\.xml$', n):
        x = z.read(n).decode('utf-8')
        if '<p:anim' in x or '<p:set' in x:
            anim += 1
        if 'p:audio' in x or '.mp3' in x or '.wav' in x or '.m4a' in x:
            audio += 1
        if 'advTm' in x or 'advClick' in x:
            print('TIMED ADVANCE FOUND in', n)
print('object_animation_slides =', anim, 'audio_slides =', audio)

pdf = 'E:/PG/2026 Hydrological cycle/sections/PROD-W3-QIANWEN-OFFICE/L16_W3_QianwenOffice_Slides36-37_39_v01.pdf'
raw = open(pdf, 'rb').read()
print('pdf_count =', re.findall(rb'/Count \d+', raw)[:3], 'page_objs =', len(re.findall(rb'/Type/Page[^s]', raw)))
print('VERIFY_DONE')
