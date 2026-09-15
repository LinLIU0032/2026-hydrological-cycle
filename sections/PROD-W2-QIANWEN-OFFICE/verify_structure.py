import sys, io, zipfile, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.util import Emu

path = 'E:/PG/2026 Hydrological cycle/sections/PROD-W2-QIANWEN-OFFICE/L16_W2_QianwenOffice_Slides16-17_v01.pptx'
prs = Presentation(path)
print('slide_count =', len(list(prs.slides)))
print('size_in = %.4f x %.4f' % (prs.slide_width/914400, prs.slide_height/914400))

def count_pictures(shape):
    n = 0
    if shape.shape_type == 13:
        n += 1
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            n += count_pictures(s)
    return n

for i, slide in enumerate(prs.slides, 1):
    pics = sum(count_pictures(sh) for sh in slide.shapes)
    note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    nshapes = len(slide.shapes)
    print(f'slide {i}: shapes={nshapes} pictures={pics} notes_chars={len(note)}')

z = zipfile.ZipFile(path)
ext = [n for n in z.namelist() if n.startswith('ppt/slides/_rels/')]
ext_targets = []
for n in ext:
    data = z.read(n).decode('utf-8')
    for m in re.finditer(r'TargetMode="External"', data):
        ext_targets.append(n)
print('external_relationships =', len(ext_targets))
media = [n for n in z.namelist() if n.startswith('ppt/media/')]
print('media =', media)
anim = 0
audio = 0
for n in z.namelist():
    if re.match(r'ppt/slides/slide\d+\.xml$', n):
        x = z.read(n).decode('utf-8')
        if '<p:anim' in x or '<p:set' in x or '<p:animEffect' in x:
            anim += 1
        if 'p:audio' in x or '.wav' in x or '.mp3' in x or '.m4a' in x:
            audio += 1
print('object_animation_slides =', anim, 'audio_slides =', audio)

import subprocess
pdf = 'E:/PG/2026 Hydrological cycle/sections/PROD-W2-QIANWEN-OFFICE/L16_W2_QianwenOffice_Slides16-17_v01.pdf'
with open(pdf, 'rb') as f:
    raw = f.read()
print('pdf_pages =', raw.count(b'/Type /Page') - raw.count(b'/Type /Pages'))
print('VERIFY_DONE')
