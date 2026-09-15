from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"E:\PG\2026 Hydrological cycle")
TARGET_SLIDES = {18, 19, 20, 28, 29, 30, 31, 32, 33, 34, 36, 37, 39, 40, 41, 42, 43}
TARGET_FIGURES = {"W-02", "W-06", "C-01", "C-04", "C-05", "C-06", "F-01", "F-02", "I-01", "I-02"}


def norm(value):
    if value is None:
        return ""
    return str(value).strip()


def numbers(value: str) -> set[int]:
    return {int(x) for x in re.findall(r"(?<!\d)(\d{1,2})(?!\d)", value)}


def inspect(path: Path, mode: str) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    result = {"path": str(path), "sheets": {}}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_idx = None
        headers = []
        for idx, row in enumerate(rows[:20]):
            values = [norm(v) for v in row]
            joined = " | ".join(values).lower()
            if (mode == "source" and "slide" in joined and "source" in joined) or (
                mode == "figure" and "figure" in joined and "status" in joined
            ):
                header_idx = idx
                headers = values
                break
        if header_idx is None:
            continue
        matched = []
        for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            values = [norm(v) for v in row]
            record = {headers[i] or f"col_{i+1}": values[i] for i in range(min(len(headers), len(values)))}
            joined = " | ".join(values)
            if mode == "source":
                slide_fields = " ".join(v for k, v in record.items() if "slide" in k.lower())
                if numbers(slide_fields) & TARGET_SLIDES:
                    matched.append({"row": excel_row, "record": record})
            else:
                figure_fields = " ".join(v for k, v in record.items() if "figure" in k.lower() or k.lower() == "id")
                if any(fid in figure_fields for fid in TARGET_FIGURES):
                    matched.append({"row": excel_row, "record": record})
        if matched:
            result["sheets"][ws.title] = {"header_row": header_idx + 1, "matches": matched}
    return result


output = {
    "source_matrix": inspect(ROOT / "planning" / "L16_SOURCE_MATRIX.xlsx", "source"),
    "figure_index": inspect(ROOT / "planning" / "L16_FIGURE_INDEX.xlsx", "figure"),
}
out_path = ROOT / "qa" / "GATE8_QA_20260914" / "planning_workbook_extract.json"
out_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
print(out_path)
