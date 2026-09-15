from __future__ import annotations

import hashlib
import io
import json
import math
import re
import zipfile
from pathlib import Path

from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(r"E:\PG\2026 Hydrological cycle")
OUT = ROOT / "qa" / "GATE8_QA_20260914" / "gate8_inspection.json"

DECKS = {
    "QODER-CN": {
        "path": ROOT / "sections" / "PROD-W3-QODER-CN" / "L16_W3_QoderCN_Slides18-20_34_v01.pptx",
        "expected_pages": [18, 19, 20, 34],
    },
    "WORKBUDDY": {
        "path": ROOT / "sections" / "PROD-W3-WORKBUDDY" / "L16_W3_WorkBuddy_Slides28-33_v01.pptx",
        "expected_pages": [28, 29, 30, 31, 32, 33],
    },
    "QIANWEN-OFFICE": {
        "path": ROOT / "sections" / "PROD-W3-QIANWEN-OFFICE" / "L16_W3_QianwenOffice_Slides36-37_39_v01.pptx",
        "expected_pages": [36, 37, 39],
    },
    "DOUBAO": {
        "path": ROOT / "sections" / "PROD-W3-DOUBAO" / "L16_W3_Doubao_Slides40-43_v01.pptx",
        "expected_pages": [40, 41, 42, 43],
    },
}

APPROVED_ASSETS = {
    "W-02": ROOT / "assets" / "figures" / "approved" / "W-02_EP_InterludeF_HydrologicCycle_pdf615-616_print580-581_spread_raw.png",
    "C-01": ROOT / "assets" / "figures" / "approved" / "C-01_UE_Fig12-19_PDF1209_raw.png",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text).strip()


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    parts = []
    for shape in slide.notes_slide.shapes:
        text = shape_text(shape)
        if text and text not in {"1", "2", "3", "4", "5", "6"}:
            parts.append(text)
    return "\n".join(parts).strip()


def slide_number_from_text(text: str, expected: int) -> bool:
    return bool(re.search(rf"(?<!\d){expected}(?!\d)", text))


def inspect_deck(name: str, spec: dict) -> dict:
    path = spec["path"]
    prs = Presentation(path)
    deck = {
        "path": str(path),
        "sha256": sha256_file(path),
        "slide_count": len(prs.slides),
        "size_emu": [prs.slide_width, prs.slide_height],
        "size_inches": [round(prs.slide_width / 914400, 4), round(prs.slide_height / 914400, 4)],
        "expected_pages": spec["expected_pages"],
        "slides": [],
    }
    for idx, (slide, expected) in enumerate(zip(prs.slides, spec["expected_pages"]), start=1):
        shapes = list(iter_shapes(slide.shapes))
        texts = [t for t in (shape_text(s) for s in shapes) if t]
        visible_text = "\n".join(texts)
        notes = notes_text(slide)
        pictures = []
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                similarity = {}
                with Image.open(io.BytesIO(image.blob)).convert("RGB") as embedded:
                    for asset_id, asset_path in APPROVED_ASSETS.items():
                        with Image.open(asset_path).convert("RGB") as approved:
                            resized = approved.resize(embedded.size, Image.Resampling.LANCZOS)
                            stat = ImageStat.Stat(ImageChops.difference(embedded, resized))
                            similarity[asset_id] = round(
                                math.sqrt(sum(v * v for v in stat.rms) / 3), 4
                            )
                pictures.append(
                    {
                        "name": shape.name,
                        "content_type": image.content_type,
                        "pixel_size": list(image.size),
                        "bytes": len(image.blob),
                        "sha256": sha256_bytes(image.blob),
                        "position_emu": [shape.left, shape.top, shape.width, shape.height],
                        "crop": [shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom],
                        "approved_asset_rms_0_255": similarity,
                    }
                )
        deck["slides"].append(
            {
                "local_index": idx,
                "expected_page": expected,
                "page_number_present": slide_number_from_text(visible_text, expected),
                "shape_count_recursive": len(shapes),
                "picture_count": len(pictures),
                "pictures": pictures,
                "visible_text": visible_text,
                "notes_length": len(notes),
                "notes_text": notes,
            }
        )

    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        external = []
        for member in names:
            if member.endswith(".rels"):
                data = zf.read(member).decode("utf-8", errors="replace")
                if 'TargetMode="External"' in data:
                    external.append(member)
        slide_xml = [zf.read(n).decode("utf-8", errors="replace") for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        deck["package"] = {
            "external_relationship_parts": external,
            "timing_slide_count": sum("<p:timing" in x for x in slide_xml),
            "timed_advance_slide_count": sum(bool(re.search(r"advTm=|advClick=", x)) for x in slide_xml),
            "audio_reference_slide_count": sum(bool(re.search(r"audio|snd|wavAudioFile", x, re.I)) for x in slide_xml),
            "media": [
                {
                    "part": n,
                    "bytes": len(zf.read(n)),
                    "sha256": sha256_bytes(zf.read(n)),
                }
                for n in names
                if n.startswith("ppt/media/") and not n.endswith("/")
            ],
        }
    return deck


result = {
    "approved_assets": {
        key: {
            "path": str(path),
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
            "pixel_size": list(Image.open(path).size),
        }
        for key, path in APPROVED_ASSETS.items()
    },
    "decks": {name: inspect_deck(name, spec) for name, spec in DECKS.items()},
}

OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
