from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(r"E:\PG\2026 Hydrological cycle")
OUT = ROOT / "qa" / "GATE8_R1_QA_20260915" / "focused_inspection.json"
PREVIOUS = ROOT / "qa" / "GATE8_QA_20260914" / "gate8_inspection.json"
DECKS = {
    "QIANWEN-OFFICE": {
        "path": ROOT / "sections" / "PROD-W3-QIANWEN-OFFICE" / "L16_W3_QianwenOffice_Slides36-37_39_v01.pptx",
        "expected_pages": [36, 37, 39],
    },
    "DOUBAO": {
        "path": ROOT / "sections" / "PROD-W3-DOUBAO" / "L16_W3_Doubao_Slides40-43_v01.pptx",
        "expected_pages": [40, 41, 42, 43],
    },
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text).strip()


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    parts = []
    for shape in slide.notes_slide.shapes:
        text = shape_text(shape)
        if text and text not in {"1", "2", "3", "4", "5", "6"}:
            parts.append(text)
    return "\n".join(parts).strip()


previous = json.loads(PREVIOUS.read_text(encoding="utf-8"))
result = {"decks": {}}

for name, spec in DECKS.items():
    path = spec["path"]
    prs = Presentation(path)
    old_slides = {
        entry["expected_page"]: entry
        for entry in previous["decks"][name]["slides"]
    }
    deck = {
        "path": str(path),
        "sha256": sha256(path),
        "previous_sha256": previous["decks"][name]["sha256"],
        "slide_count": len(prs.slides),
        "size_inches": [round(prs.slide_width / 914400, 4), round(prs.slide_height / 914400, 4)],
        "slides": [],
    }
    for local_index, (slide, expected_page) in enumerate(zip(prs.slides, spec["expected_pages"]), start=1):
        visible = "\n".join(
            text for text in (shape_text(shape) for shape in iter_shapes(slide.shapes)) if text
        )
        notes = notes_text(slide)
        old = old_slides[expected_page]
        deck["slides"].append(
            {
                "local_index": local_index,
                "expected_page": expected_page,
                "page_number_present": bool(re.search(rf"(?<!\d){expected_page}(?!\d)", visible)),
                "visible_text": visible,
                "notes_text": notes,
                "visible_text_unchanged": visible == old["visible_text"],
                "notes_text_unchanged": notes == old["notes_text"],
            }
        )

    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        external = []
        for member in names:
            if member.endswith(".rels"):
                rels = zf.read(member).decode("utf-8", errors="replace")
                if 'TargetMode="External"' in rels:
                    external.append(member)
        slide_xml = [
            zf.read(member).decode("utf-8", errors="replace")
            for member in names
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", member)
        ]
        deck["package"] = {
            "notes_parts": sum(bool(re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", member)) for member in names),
            "external_relationship_parts": external,
            "timing_slide_count": sum("<p:timing" in xml for xml in slide_xml),
            "timed_advance_slide_count": sum(bool(re.search(r"advTm=|advClick=", xml)) for xml in slide_xml),
            "audio_reference_slide_count": sum(bool(re.search(r"audio|snd|wavAudioFile", xml, re.I)) for xml in slide_xml),
        }
    result["decks"][name] = deck

slides = {
    slide["expected_page"]: slide
    for deck in result["decks"].values()
    for slide in deck["slides"]
}
result["focused_checks"] = {
    "slide_36_forbidden_values_removed": not any(
        token in slides[36][field]
        for token in ("7%/°C", "75%", "百分之七", "百分之七十五")
        for field in ("visible_text", "notes_text")
    ),
    "slide_37_visible_meta_removed": not any(
        token in slides[37]["visible_text"]
        for token in ("No ice-volume", "sea-level", "Snowball-Earth")
    ),
    "slide_39_incomplete_phrase_removed": "植物生长可促进" not in (
        slides[39]["visible_text"] + slides[39]["notes_text"]
    ),
    "slide_40_plant_time_range_removed": "生长季—年代际" not in (
        slides[40]["visible_text"] + slides[40]["notes_text"]
    ),
    "slide_40_keeps_no_unified_clock": "响应时钟未统一量化" in slides[40]["visible_text"],
    "slide_43_exclusive_weathering_definition_removed": "风化在这里指硅酸盐化学风化对二氧化碳的去除" not in slides[43]["notes_text"],
    "slide_43_keeps_weathering_distinction": all(
        token in slides[43]["notes_text"] for token in ("硅酸盐风化", "碳酸盐风化")
    ),
    "slide_41_visible_regression": slides[41]["visible_text_unchanged"],
    "slide_41_notes_regression": slides[41]["notes_text_unchanged"],
    "slide_42_visible_regression": slides[42]["visible_text_unchanged"],
    "slide_42_notes_regression": slides[42]["notes_text_unchanged"],
}

OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
