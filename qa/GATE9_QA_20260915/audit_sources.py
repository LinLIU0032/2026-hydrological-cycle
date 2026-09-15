from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(r"E:\PG\2026 Hydrological cycle")
OUT = ROOT / "qa" / "GATE9_QA_20260915" / "accepted_source_manifest.json"

DECKS = {
    "SAMPLE": ("review/L16_STYLE_SAMPLE_v01.pptx", 9),
    "W1_QODER": ("sections/PROD-W1-QODER-CN/L16_W1_QoderCN_Slides04-06_v01.pptx", 3),
    "W2_QODER": ("sections/PROD-W2-QODER-CN/L16_W2_QoderCN_Slides07-08_v01.pptx", 2),
    "W1_DOUBAO": ("sections/PROD-W1-DOUBAO/L16_W1_Doubao_Slides10-12_v01.pptx", 3),
    "W2_DOUBAO": ("sections/PROD-W2-DOUBAO/L16_W2_Doubao_Slides13-15_v01.pptx", 3),
    "W2_QIANWEN": ("sections/PROD-W2-QIANWEN-OFFICE/L16_W2_QianwenOffice_Slides16-17_v01.pptx", 2),
    "W3_QODER": ("sections/PROD-W3-QODER-CN/L16_W3_QoderCN_Slides18-20_34_v01.pptx", 4),
    "W1_WORKBUDDY": ("sections/PROD-W1-WORKBUDDY/L16_W1_WorkBuddy_Slides21-22-24_v01.pptx", 3),
    "W2_WORKBUDDY": ("sections/PROD-W2-WORKBUDDY/L16_W2_WorkBuddy_Slides25-26_v01.pptx", 2),
    "W3_WORKBUDDY": ("sections/PROD-W3-WORKBUDDY/L16_W3_WorkBuddy_Slides28-33_v01.pptx", 6),
    "W3_QIANWEN": ("sections/PROD-W3-QIANWEN-OFFICE/L16_W3_QianwenOffice_Slides36-37_39_v01.pptx", 3),
    "W3_DOUBAO": ("sections/PROD-W3-DOUBAO/L16_W3_Doubao_Slides40-43_v01.pptx", 4),
}

MAPPING = [
    (1, "SAMPLE", 1), (2, "SAMPLE", 2), (3, "SAMPLE", 3),
    (4, "W1_QODER", 1), (5, "W1_QODER", 2), (6, "W1_QODER", 3),
    (7, "W2_QODER", 1), (8, "W2_QODER", 2), (9, "SAMPLE", 4),
    (10, "W1_DOUBAO", 1), (11, "W1_DOUBAO", 2), (12, "W1_DOUBAO", 3),
    (13, "W2_DOUBAO", 1), (14, "W2_DOUBAO", 2), (15, "W2_DOUBAO", 3),
    (16, "W2_QIANWEN", 1), (17, "W2_QIANWEN", 2),
    (18, "W3_QODER", 1), (19, "W3_QODER", 2), (20, "W3_QODER", 3),
    (21, "W1_WORKBUDDY", 1), (22, "W1_WORKBUDDY", 2), (23, "SAMPLE", 5),
    (24, "W1_WORKBUDDY", 3), (25, "W2_WORKBUDDY", 1), (26, "W2_WORKBUDDY", 2),
    (27, "SAMPLE", 6),
    (28, "W3_WORKBUDDY", 1), (29, "W3_WORKBUDDY", 2), (30, "W3_WORKBUDDY", 3),
    (31, "W3_WORKBUDDY", 4), (32, "W3_WORKBUDDY", 5), (33, "W3_WORKBUDDY", 6),
    (34, "W3_QODER", 4), (35, "SAMPLE", 7),
    (36, "W3_QIANWEN", 1), (37, "W3_QIANWEN", 2), (38, "SAMPLE", 8),
    (39, "W3_QIANWEN", 3),
    (40, "W3_DOUBAO", 1), (41, "W3_DOUBAO", 2), (42, "W3_DOUBAO", 3),
    (43, "W3_DOUBAO", 4), (44, "SAMPLE", 9),
]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text).strip()


def slide_text(slide) -> str:
    return "\n".join(text for text in (shape_text(shape) for shape in iter_shapes(slide.shapes)) if text)


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return "\n".join(
        text for text in (shape_text(shape) for shape in slide.notes_slide.shapes)
        if text and text not in {"1", "2", "3", "4", "5", "6"}
    )


errors = []
decks = {}
presentations = {}
for deck_id, (relative_path, expected_count) in DECKS.items():
    path = ROOT / relative_path
    if not path.exists():
        errors.append(f"missing source: {relative_path}")
        continue
    prs = Presentation(path)
    presentations[deck_id] = prs
    if len(prs.slides) != expected_count:
        errors.append(f"{deck_id}: expected {expected_count} slides, found {len(prs.slides)}")
    if (prs.slide_width, prs.slide_height) != (12192000, 6858000):
        errors.append(f"{deck_id}: unexpected canvas {prs.slide_width}x{prs.slide_height}")
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        notes_parts = sum(bool(re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)) for name in names)
        external = []
        for name in names:
            if name.endswith(".rels") and 'TargetMode="External"' in zf.read(name).decode("utf-8", errors="replace"):
                external.append(name)
        if notes_parts != expected_count:
            errors.append(f"{deck_id}: expected {expected_count} notes parts, found {notes_parts}")
        if external:
            errors.append(f"{deck_id}: external relationships in {external}")
        decks[deck_id] = {
            "path": str(path),
            "sha256": sha256_bytes(path.read_bytes()),
            "bytes": path.stat().st_size,
            "slide_count": len(prs.slides),
            "notes_parts": notes_parts,
            "size_emu": [prs.slide_width, prs.slide_height],
            "external_relationship_parts": external,
        }

pages = []
if [page for page, _, _ in MAPPING] != list(range(1, 45)):
    errors.append("final page roster is not exactly 1..44")
for final_page, deck_id, local_slide in MAPPING:
    prs = presentations.get(deck_id)
    if prs is None or local_slide < 1 or local_slide > len(prs.slides):
        errors.append(f"page {final_page}: invalid mapping {deck_id} slide {local_slide}")
        continue
    slide = prs.slides[local_slide - 1]
    visible = slide_text(slide)
    notes = notes_text(slide)
    expected_folio = f"{final_page:02d}"
    if final_page != 1 and not re.search(rf"(?<!\d){re.escape(expected_folio)}(?!\d)", visible):
        errors.append(f"page {final_page}: visible page number {expected_folio} not found")
    if not notes:
        errors.append(f"page {final_page}: empty notes")
    pages.append(
        {
            "final_page": final_page,
            "deck_id": deck_id,
            "source_path": decks[deck_id]["path"],
            "local_slide": local_slide,
            "visible_text_sha256": sha256_bytes(visible.encode("utf-8")),
            "notes_text_sha256": sha256_bytes(notes.encode("utf-8")),
            "visible_text": visible,
            "notes_text": notes,
        }
    )

result = {"status": "passed" if not errors else "failed", "errors": errors, "decks": decks, "pages": pages}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
if errors:
    raise SystemExit("\n".join(errors))
