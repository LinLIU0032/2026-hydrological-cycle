from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(r"E:\PG\2026 Hydrological cycle")
QA_DIR = ROOT / "qa" / "GATE9_QA_20260915"
MASTER = ROOT / "final" / "L16_Hydrological_Carbon_Cycles_v02.pptx"
SOURCE_MANIFEST = QA_DIR / "accepted_source_manifest.json"
DELIVERY_CHECKER = (
    ROOT
    / "planning"
    / "tooling"
    / "ppt-master-v6.3.1"
    / "ppt-master"
    / "skills"
    / "ppt-master"
    / "scripts"
    / "pptx_delivery_check.py"
)
OUT = QA_DIR / "master_structural_audit_v02.json"
DELIVERY_OUT = QA_DIR / "master_delivery_check_v02.json"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text).strip()


def slide_text(slide) -> str:
    return "\n".join(
        text for text in (shape_text(shape) for shape in iter_shapes(slide.shapes)) if text
    )


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return "\n".join(
        text
        for text in (shape_text(shape) for shape in slide.notes_slide.shapes)
        if text and text not in {"1", "2", "3", "4", "5", "6"}
    )


def local_name(name: str) -> str:
    return name.rsplit("}", 1)[-1]


def descendants_by_local_name(element, name: str):
    return [node for node in element.iter() if local_name(node.tag) == name]


def transition_semantic(slide_element):
    candidates = descendants_by_local_name(slide_element, "transition")
    if not candidates:
        return None
    transition = max(
        candidates,
        key=lambda node: int(any(local_name(key) == "dur" for key in node.attrib)),
    )
    attrs = {local_name(key): value for key, value in transition.attrib.items()}
    effects = [local_name(child.tag) for child in transition]
    return {
        "effect": effects[0] if effects else None,
        "duration_ms": attrs.get("dur"),
        "speed": attrs.get("spd"),
        "advance_on_click": attrs.get("advClick"),
        "advance_time_ms": attrs.get("advTm"),
    }


source = json.loads(SOURCE_MANIFEST.read_text(encoding="utf-8"))
errors: list[str] = []
warnings: list[str] = []
if source.get("status") != "passed":
    errors.append("accepted source manifest is not passed")

delivery_process = subprocess.run(
    [sys.executable, str(DELIVERY_CHECKER), str(MASTER)],
    check=False,
    capture_output=True,
    text=True,
    encoding="utf-8",
)
if delivery_process.returncode != 0:
    errors.append(f"delivery checker exited {delivery_process.returncode}: {delivery_process.stderr.strip()}")
    delivery = None
else:
    delivery = json.loads(delivery_process.stdout)
    DELIVERY_OUT.write_text(json.dumps(delivery, ensure_ascii=False, indent=2), encoding="utf-8")
    if delivery.get("status") not in {"passed", "passed-with-advisories"}:
        errors.append(f"delivery checker status: {delivery.get('status')}")
    warnings.extend(delivery.get("advisories", []))

master = Presentation(MASTER)
if len(master.slides) != 44:
    errors.append(f"expected 44 slides, found {len(master.slides)}")
if (master.slide_width, master.slide_height) != (12192000, 6858000):
    errors.append(f"unexpected canvas {master.slide_width}x{master.slide_height}")

source_decks = {
    deck_id: Presentation(deck["path"])
    for deck_id, deck in source["decks"].items()
}
page_results = []
for expected, final_slide in zip(source["pages"], master.slides):
    final_page = expected["final_page"]
    visible = slide_text(final_slide)
    notes = notes_text(final_slide)
    visible_hash = sha256_bytes(visible.encode("utf-8"))
    notes_hash = sha256_bytes(notes.encode("utf-8"))
    visible_match = visible_hash == expected["visible_text_sha256"]
    notes_match = notes_hash == expected["notes_text_sha256"]
    if not visible_match:
        errors.append(f"page {final_page}: visible text differs from accepted source")
    if not notes_match:
        errors.append(f"page {final_page}: Notes differ from accepted source")
    if final_page != 1 and f"{final_page:02d}" not in visible:
        errors.append(f"page {final_page}: zero-padded folio not found")

    source_slide = source_decks[expected["deck_id"]].slides[expected["local_slide"] - 1]
    source_transition = transition_semantic(source_slide._element)
    final_transition = transition_semantic(final_slide._element)
    transition_match = source_transition == final_transition
    if not transition_match:
        errors.append(f"page {final_page}: transition differs from accepted source")
    if descendants_by_local_name(final_slide._element, "timing"):
        errors.append(f"page {final_page}: object-animation timing tree found")
    transition_elements = descendants_by_local_name(final_slide._element, "transition")
    if any(any(local_name(key) == "advTm" for key in node.attrib) for node in transition_elements):
        errors.append(f"page {final_page}: timed advance found")

    page_results.append(
        {
            "final_page": final_page,
            "deck_id": expected["deck_id"],
            "local_slide": expected["local_slide"],
            "visible_text_match": visible_match,
            "notes_text_match": notes_match,
            "transition_match": transition_match,
        }
    )

with zipfile.ZipFile(MASTER) as zf:
    names = zf.namelist()
    notes_parts = [name for name in names if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")]
    if len(notes_parts) != 44:
        errors.append(f"expected 44 notes parts, found {len(notes_parts)}")
    external_relationship_parts = []
    for name in names:
        if name.endswith(".rels") and b'TargetMode="External"' in zf.read(name):
            external_relationship_parts.append(name)
    if external_relationship_parts:
        errors.append(f"external relationships found: {external_relationship_parts}")
    media_parts = [name for name in names if name.startswith("ppt/media/")]
    forbidden_media = [
        name
        for name in media_parts
        if Path(name).suffix.lower() in {".mp3", ".m4a", ".wav", ".wma", ".mp4", ".mov", ".avi", ".wmv"}
    ]
    if forbidden_media:
        errors.append(f"audio/video media found: {forbidden_media}")
    final_media_hashes = {sha256_bytes(zf.read(name)) for name in media_parts}

accepted_media_hashes = set()
for deck in source["decks"].values():
    with zipfile.ZipFile(deck["path"]) as zf:
        for name in zf.namelist():
            if name.startswith("ppt/media/"):
                accepted_media_hashes.add(sha256_bytes(zf.read(name)))
unapproved_media_hashes = sorted(final_media_hashes - accepted_media_hashes)
if unapproved_media_hashes:
    errors.append(f"master contains {len(unapproved_media_hashes)} media hashes not present in accepted sources")

result = {
    "status": "passed" if not errors else "failed",
    "master": {
        "path": str(MASTER),
        "sha256": sha256_bytes(MASTER.read_bytes()),
        "bytes": MASTER.stat().st_size,
        "slide_count": len(master.slides),
        "notes_parts": len(notes_parts),
        "size_emu": [master.slide_width, master.slide_height],
        "media_parts": len(media_parts),
        "unique_media_hashes": len(final_media_hashes),
        "external_relationship_parts": external_relationship_parts,
    },
    "delivery_checker_status": None if delivery is None else delivery.get("status"),
    "page_results": page_results,
    "warnings": warnings,
    "errors": errors,
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
if errors:
    raise SystemExit("\n".join(errors))
